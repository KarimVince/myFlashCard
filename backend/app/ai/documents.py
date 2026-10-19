"""Text from a member's document (PDF, Word, plain text), used as source material for AI generation.

Only the text is kept, and only for the request: the file itself is never stored.
"""
import io
import re

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_FILE_BYTES = 5 * 1024 * 1024
# About 3,000 tokens: leaves room for the instructions and the deck within free-tier request limits.
MAX_SOURCE_CHARS = 12_000
MAX_PDF_PAGES = 60

ACCEPTED = {".pdf", ".docx", ".txt", ".md", ".csv"}


class DocumentError(ValueError):
    """The file can't be used. The message is shown to the member."""


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def _pdf(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise DocumentError("This PDF is password-protected. Remove the password and try again.")
        pages = reader.pages[:MAX_PDF_PAGES]
        return "\n\n".join(page.extract_text() or "" for page in pages)
    except PdfReadError:
        raise DocumentError("This PDF can't be read. Try saving it again or exporting it as text.")


def _docx(content: bytes) -> str:
    try:
        doc = Document(io.BytesIO(content))
    except Exception:
        raise DocumentError("This Word file can't be read. Save it as .docx and try again.")
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _plain(content: bytes) -> str:
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentError("This text file can't be read.")


def extract_text(filename: str, content: bytes) -> tuple[str, bool]:
    """Return (text, truncated). Raises DocumentError with a message for the member."""
    name = (filename or "").lower()
    ext = name[name.rfind("."):] if "." in name else ""
    if ext not in ACCEPTED:
        raise DocumentError("Use a PDF, Word (.docx), or text file (.txt, .md, .csv).")
    if len(content) > MAX_FILE_BYTES:
        raise DocumentError("The file is larger than 5 MB.")
    if ext == ".pdf":
        raw = _pdf(content)
    elif ext == ".docx":
        raw = _docx(content)
    else:
        raw = _plain(content)
    text = _clean(raw)
    if len(text) < 20:
        raise DocumentError(
            "No text found in this file. Scanned PDFs (pictures of pages) aren't supported yet."
        )
    if len(text) > MAX_SOURCE_CHARS:
        return text[:MAX_SOURCE_CHARS].rstrip(), True
    return text, False
